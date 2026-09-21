# -*- coding: utf-8 -*-
"""Buduje kompletny szablon referencyjny (reference_verdana.docx) od zera,
w jednym spojnym skrypcie, laczac wszystkie wczesniejsze poprawki:
Verdana, justowanie, style niestandardowe, stopka z numeracja stron,
marginesy A4, styl tabeli bez obramowan i wysrodkowany.

UWAGA (Zenek/Claude Code): to jest CELOWO odtwarzanie szablonu OD ZERA
kodem, a nie poleganie na gotowym pliku binarnym reference_verdana.docx -
patrz README.md w tym katalogu: w oryginalnym srodowisku klientki (Projekt
na claude.ai bez uruchamiania kodu) wgrany plik .docx nie przetrwal jako
prawdziwy plik binarny, wiec jedynym niezawodnym sposobem odtworzenia
szablonu jest ten skrypt. Uruchamiac go raz na poczatku pracy nad nowym
zleceniem (lub gdy assets/reference_verdana.docx nie istnieje), zanim
odpali sie build2.py."""
import os, re, shutil, subprocess
import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, Mm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = "assets/reference_base.docx"
UNPACK = "assets/reference_build_unpacked"
OUT = "assets/reference_verdana.docx"

# 1) Domyslny szablon pandoc
subprocess.run(["pandoc", "-o", BASE, "--print-default-data-file", "reference.docx"], check=True)

# 2) Rozpakuj i podmien czcionki (Calibri/Cambria -> Verdana) w motywie i w styles.xml
if os.path.exists(UNPACK):
    shutil.rmtree(UNPACK)
os.makedirs(UNPACK)
with __import__("zipfile").ZipFile(BASE) as z:
    z.extractall(UNPACK)

styles_path = os.path.join(UNPACK, "word", "styles.xml")
with open(styles_path, encoding="utf-8") as f:
    xml = f.read()
xml = re.sub(r'(w:(?:ascii|hAnsi|cs|eastAsia)=")[^"]*(")', r'\1Verdana\2', xml)
with open(styles_path, "w", encoding="utf-8") as f:
    f.write(xml)

theme_path = os.path.join(UNPACK, "word", "theme", "theme1.xml")
with open(theme_path, encoding="utf-8") as f:
    xml = f.read()
xml = xml.replace('typeface="Calibri"', 'typeface="Verdana"').replace('typeface="Cambria"', 'typeface="Verdana"')
with open(theme_path, "w", encoding="utf-8") as f:
    f.write(xml)

# 3) Styl "Table": wysrodkowany, bez obramowan (uzywany tylko do rzedu logotypow)
styles_path2 = os.path.join(UNPACK, "word", "styles.xml")
with open(styles_path2, encoding="utf-8") as f:
    xml = f.read()
new_table_style = (
    '<w:style w:type="table" w:default="1" w:styleId="Table"><w:name w:val="Table"/>'
    '<w:basedOn w:val="TableNormal"/><w:semiHidden/><w:unhideWhenUsed/><w:qFormat/>'
    '<w:tblPr><w:jc w:val="center"/><w:tblInd w:w="0" w:type="dxa"/>'
    '<w:tblBorders>'
    '<w:top w:val="none" w:sz="0" w:space="0"/><w:left w:val="none" w:sz="0" w:space="0"/>'
    '<w:bottom w:val="none" w:sz="0" w:space="0"/><w:right w:val="none" w:sz="0" w:space="0"/>'
    '<w:insideH w:val="none" w:sz="0" w:space="0"/><w:insideV w:val="none" w:sz="0" w:space="0"/>'
    '</w:tblBorders>'
    '<w:tblCellMar><w:top w:w="0" w:type="dxa"/><w:left w:w="108" w:type="dxa"/>'
    '<w:bottom w:w="0" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar></w:tblPr>'
    '<w:tblStylePr w:type="firstRow"><w:tblPr><w:jc w:val="center"/><w:tblInd w:w="0" w:type="dxa"/></w:tblPr>'
    '<w:trPr><w:jc w:val="center"/></w:trPr>'
    '<w:tcPr><w:vAlign w:val="center"/><w:tcBorders>'
    '<w:top w:val="none" w:sz="0"/><w:left w:val="none" w:sz="0"/>'
    '<w:bottom w:val="none" w:sz="0"/><w:right w:val="none" w:sz="0"/>'
    '</w:tcBorders></w:tcPr></w:tblStylePr></w:style>'
)
xml2, n = re.subn(r'<w:style [^>]*w:styleId="Table"[^>]*>.*?</w:style>', new_table_style, xml, flags=re.DOTALL)
assert n == 1
with open(styles_path2, "w", encoding="utf-8") as f:
    f.write(xml2)

# 3b) Domyslny rozmiar czcionki w calym dokumencie: 12pt -> 10pt
# UWAGA: poprzednia wersja uzywala literalnego str.replace() dopasowanego do
# "<w:sz.../><w:szCs.../>" bez spacji przed "/>". Domyslny reference.docx
# pandoca zapisuje te znaczniki z odstepem (`<w:sz w:val="24" />`) i w osobnych
# liniach, wiec ten replace NIGDY sie nie wykonywal (0 dopasowan) - to byla
# przyczyna zrodlowa bledu "za duza czcionka" (tekst wychodzil 12pt zamiast
# 10pt, bo <w:docDefaults> nadal mial w:sz=24). Naprawione ponizej regexem
# ograniczonym wylacznie do bloku <w:docDefaults>...</w:docDefaults>.
with open(styles_path2, encoding="utf-8") as f:
    xml = f.read()

def _fix_docdefaults_size(m):
    block = m.group(0)
    block = re.sub(r'<w:sz w:val="24"\s*/>', '<w:sz w:val="20"/>', block)
    block = re.sub(r'<w:szCs w:val="24"\s*/>', '<w:szCs w:val="20"/>', block)
    return block

xml, _n = re.subn(r'<w:docDefaults>.*?</w:docDefaults>', _fix_docdefaults_size, xml, flags=re.DOTALL)
assert _n == 1, "Nie znaleziono bloku <w:docDefaults> - sprawdz strukture domyslnego reference.docx pandoca"
with open(styles_path2, "w", encoding="utf-8") as f:
    f.write(xml)

# 4) Spakuj posredni plik
tmp_docx = "assets/reference_stage1.docx"
if os.path.exists(tmp_docx):
    os.remove(tmp_docx)
shutil.make_archive("assets/reference_stage1", "zip", UNPACK)
os.rename("assets/reference_stage1.zip", tmp_docx)

# 5) Reszta poprawek przez python-docx
d = docx.Document(tmp_docx)
styles = d.styles

for name in ["Normal", "ListParagraph", "BodyText", "Compact", "FirstParagraph"]:
    try:
        st = styles[name]
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.widow_control = False
    except KeyError:
        pass

# Zabezpieczenie: wymuszamy Verdana 10pt wprost na stylu "Normal", niezaleznie
# od tego czy podmiana w <w:docDefaults> (krok 3b) sie powiodla. Wszystkie inne
# style tekstowe (BodyText, Compact, Justify...) dziedzicza rozmiar z "Normal",
# wiec to jedno ustawienie gwarantuje spojne 10pt w calym dokumencie.
styles["Normal"].font.name = "Verdana"
styles["Normal"].font.size = Pt(10)

styles["Compact"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
styles["Compact"].font.name = "Verdana"

def add_or_get(name, base=None, size=None, bold=None, italic=None, align=None, space_after=None, space_before=None):
    try:
        st = styles[name]
    except KeyError:
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if base is not None:
            st.base_style = styles[base]
    st.font.name = "Verdana"
    if size:
        st.font.size = Pt(size)
    if bold is not None:
        st.font.bold = bold
    if italic is not None:
        st.font.italic = italic
    if align is not None:
        st.paragraph_format.alignment = align
    if space_after is not None:
        st.paragraph_format.space_after = Pt(space_after)
    if space_before is not None:
        st.paragraph_format.space_before = Pt(space_before)
    return st

add_or_get("CoverCenter", base="Normal", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
add_or_get("CoverTitle", base="Normal", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
add_or_get("CoverSubtitle", base="Normal", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_or_get("CoverSmall", base="Normal", size=11, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
# UWAGA: styl "TocEntry" w prawdziwym wzorze deklaruje 11pt, ALE w samym
# dokumencie kazdy wpis spisu tresci ma bezposrednie formatowanie nadpisujace
# rozmiar na 10pt (w:sz="20") - to ten efektywny, widoczny rozmiar. Ustawiamy
# tu wprost 10pt, zeby spis tresci wygladal jak w oryginale (klientka
# zglosila: "spis treści zmniejsz czcionkę jak we wzorze moim").
add_or_get("TocEntry", base="Normal", size=10, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
add_or_get("Justify", base="Normal", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_or_get("FormText", base="Normal", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_or_get("FormHeaderCenter", base="Normal", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10, space_before=6)
add_or_get("FormRefNote", base="Normal", size=9, italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10)
add_or_get("ImgCenter", base="Normal", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, space_before=6)

# Heading1: czarny, WERSALIKI (auto-caps), wiekszy odstep pod spodem (zeby nie bylo "nazdziubolone")
h1 = styles["Heading1"]
h1.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
h1.paragraph_format.space_after = Pt(18)
h1.paragraph_format.space_before = Pt(6)
h1.element.get_or_add_rPr().append(OxmlElement('w:caps'))

# Heading2: czarny, podkreslony, wysrodkowany
h2 = styles["Heading2"]
h2.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
h2.font.underline = True
h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
h2.paragraph_format.space_before = Pt(14)
h2.paragraph_format.space_after = Pt(10)

# SectionBar: CZARNY tekst, pogrubiony, WERSALIKI, wyjustowany, na turkusowym
# tle - tytuly poszczegolnych instrukcji sluzbowych w Rozdziale V.
# Kolor tla: 4BACC6 (dokladny odcien turkusu z prawdziwego, zaakceptowanego
# wzoru - w document.xml jest to bezposredni w:shd na akapicie stylu
# "Nagwek2" z themeFill="accent5", ktorego dokladna wartosc to 4BACC6 wg
# theme1.xml). NIE 3B9AB4 - to byla wczesniejsza, bledna probka koloru.
# KOLOR TEKSTU: potwierdzone wprost przez klientke - CZARNY, nie bialy.
try:
    bar = styles["SectionBar"]
except KeyError:
    bar = styles.add_style("SectionBar", WD_STYLE_TYPE.PARAGRAPH)
    bar.base_style = styles["Normal"]
bar.font.name = "Verdana"
bar.font.bold = True
bar.font.color.rgb = docx.shared.RGBColor(0x00, 0x00, 0x00)
bar.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
bar.paragraph_format.space_before = Pt(16)
bar.paragraph_format.space_after = Pt(12)
bar.element.get_or_add_rPr().append(OxmlElement('w:caps'))
bar_ppr = bar.element.get_or_add_pPr()
shd = OxmlElement('w:shd')
shd.set(qn('w:val'), 'clear')
shd.set(qn('w:color'), 'auto')
shd.set(qn('w:fill'), '4BACC6')
bar_ppr.append(shd)

# Marginesy / rozmiar strony A4
s = d.sections[0]
s.page_width = Mm(210)
s.page_height = Mm(297)
s.top_margin = Cm(2.2)
s.bottom_margin = Cm(2.0)
s.left_margin = Cm(2.5)
s.right_margin = Cm(2.5)
s.header_distance = Cm(1.25)
s.footer_distance = Cm(1.1)

# Stopka z numeracja stron
def set_rpr(run, size=9):
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Verdana')
    rFonts.set(qn('w:hAnsi'), 'Verdana')
    rFonts.set(qn('w:cs'), 'Verdana')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(size * 2))
    rPr.append(sz)
    run._r.insert(0, rPr)

def add_field(paragraph, field_code, size=9):
    run = paragraph.add_run()
    set_rpr(run, size)
    fld_begin = OxmlElement('w:fldChar'); fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = field_code
    fld_sep = OxmlElement('w:fldChar'); fld_sep.set(qn('w:fldCharType'), 'separate')
    fld_end = OxmlElement('w:fldChar'); fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin); run._r.append(instr); run._r.append(fld_sep); run._r.append(fld_end)

footer = s.footer
footer.is_linked_to_previous = False
p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
for r in list(p.runs):
    r._r.getparent().remove(r._r)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run1 = p.add_run("Strona "); run1.font.name = "Verdana"; run1.font.size = Pt(9)
add_field(p, "PAGE")
run2 = p.add_run(" z "); run2.font.name = "Verdana"; run2.font.size = Pt(9)
add_field(p, "NUMPAGES")

d.save(OUT)
print("Zbudowano kompletny reference_verdana.docx")

# Weryfikacja
d2 = docx.Document(OUT)
names = [st.name for st in d2.styles]
for req in ["FormText", "ImgCenter", "CoverTitle", "Justify", "TocEntry", "FormHeaderCenter", "FormRefNote"]:
    print(req, "OK" if req in names else "BRAK!!!")
print("top_margin", d2.sections[0].top_margin)
print("footer text runs:", len(d2.sections[0].footer.paragraphs[0].runs))
